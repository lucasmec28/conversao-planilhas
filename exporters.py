from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Dict, Iterable, List, Sequence

import xlsxwriter
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from core import CaseMapping, ProcessedData


DARK = "#3B3B3B"
PINK = "#E98AAD"
GREEN = "#E2F0D9"
GREEN_TXT = "#385723"
RED = "#FCE4D6"
RED_TXT = "#9C0006"
LIGHT_GRAY = "#F2F2F2"
FONT_NAME = "Times New Roman"


def _status_level(processed: ProcessedData) -> str:
    levels = {v.level for v in processed.validations}
    if "error" in levels:
        return "COM ERROS"
    if "warning" in levels:
        return "COM ALERTAS"
    return "OK"


def _same_value(a, b) -> bool:
    """Comparison used only for presentation merges."""
    if a is None and b is None:
        return True
    return str(a) == str(b)


def _merge_runs(rows: Sequence[Sequence[object]], col: int, parent_cols: Sequence[int] = ()) -> List[tuple[int, int, object]]:
    """Return inclusive row-index spans for sequential equal categorical values.

    parent_cols constrains the merge so a child field (e.g. Grupo) never crosses
    the boundary of a parent field (e.g. Base/nó).
    """
    spans: List[tuple[int, int, object]] = []
    if not rows:
        return spans

    start = 0
    for i in range(1, len(rows) + 1):
        same = False
        if i < len(rows):
            same = _same_value(rows[i][col], rows[i - 1][col])
            if same and parent_cols:
                same = all(_same_value(rows[i][pc], rows[i - 1][pc]) for pc in parent_cols)
        if not same:
            if i - start >= 2:
                spans.append((start, i - 1, rows[start][col]))
            start = i
    return spans


def _scale_widths(widths_cm: Sequence[float], max_total_cm: float = 18.8) -> List[float]:
    total = float(sum(widths_cm))
    if total <= 0:
        return list(widths_cm)
    if total <= max_total_cm:
        return list(widths_cm)
    factor = max_total_cm / total
    return [w * factor for w in widths_cm]


def build_excel(
    processed: ProcessedData,
    mappings: Sequence[CaseMapping],
    input_names: Dict[str, str],
    include_raw_data: bool = True,
) -> bytes:
    bio = BytesIO()
    wb = xlsxwriter.Workbook(bio, {"in_memory": True})
    wb.set_properties({
        "title": "Tabelas Robot para Memória de Cálculo",
        "subject": processed.summary.get("structure_name", ""),
        "author": "Blossom Consult",
        "comments": "Gerado automaticamente a partir de resultados do Robot Structural Analysis.",
    })

    fmt_title = wb.add_format({
        "font_name": FONT_NAME, "font_size": 12, "bold": True,
        "align": "center", "valign": "vcenter",
    })
    fmt_header = wb.add_format({
        "font_name": FONT_NAME, "font_size": 9, "bold": True,
        "font_color": PINK, "bg_color": DARK,
        "align": "center", "valign": "vcenter", "text_wrap": True,
        "border": 1,
    })
    fmt_body = wb.add_format({
        "font_name": FONT_NAME, "font_size": 9,
        "align": "center", "valign": "vcenter", "border": 1,
    })
    fmt_body_left = wb.add_format({
        "font_name": FONT_NAME, "font_size": 9,
        "align": "left", "valign": "vcenter", "border": 1,
    })
    fmt_num = wb.add_format({
        "font_name": FONT_NAME, "font_size": 9,
        "align": "center", "valign": "vcenter", "border": 1,
        "num_format": "0.00",
    })
    fmt_ok = wb.add_format({
        "font_name": FONT_NAME, "font_size": 9, "bold": True,
        "font_color": GREEN_TXT, "bg_color": GREEN,
        "align": "center", "valign": "vcenter", "border": 1,
    })
    fmt_bad = wb.add_format({
        "font_name": FONT_NAME, "font_size": 9, "bold": True,
        "font_color": RED_TXT, "bg_color": RED,
        "align": "center", "valign": "vcenter", "border": 1,
    })
    fmt_section = wb.add_format({
        "font_name": FONT_NAME, "font_size": 10, "bold": True,
        "font_color": PINK, "bg_color": DARK, "border": 1,
        "align": "left", "valign": "vcenter",
    })
    fmt_label = wb.add_format({"font_name": FONT_NAME, "font_size": 9, "bold": True})
    fmt_text = wb.add_format({"font_name": FONT_NAME, "font_size": 9})

    # RESUMO
    ws = wb.add_worksheet("RESUMO")
    ws.hide_gridlines(2)
    ws.set_column("A:A", 28)
    ws.set_column("B:B", 22)
    ws.set_column("C:C", 36)
    ws.set_column("D:D", 16)
    ws.set_column("E:E", 20)
    ws.merge_range("A1:B1", "GERADOR DE TABELAS - ROBOT STRUCTURAL ANALYSIS", fmt_title)
    ws.write("A3", "Estrutura", fmt_label)
    ws.write("B3", processed.summary.get("structure_name", ""), fmt_text)
    ws.write("A4", "Gerado em", fmt_label)
    ws.write("B4", datetime.now().strftime("%d/%m/%Y %H:%M"), fmt_text)
    ws.write("A5", "Status", fmt_label)
    ws.write("B5", _status_level(processed), fmt_text)
    ws.write("A7", "ARQUIVOS DE ENTRADA", fmt_section)
    ws.write("A8", "Combinações", fmt_label); ws.write("B8", input_names.get("combinations", ""), fmt_text)
    ws.write("A9", "Reações", fmt_label); ws.write("B9", input_names.get("reactions", ""), fmt_text)
    ws.write("A10", "ELU/ELS", fmt_label); ws.write("B10", input_names.get("members", ""), fmt_text)
    ws.write("A12", "CONFIGURAÇÕES", fmt_section)
    ws.write("A13", "λ máximo padrão", fmt_label); ws.write("B13", processed.summary.get("lambda_limit"), fmt_text)
    ws.write("A14", "Fator kgf → kN", fmt_label); ws.write("B14", processed.summary.get("force_conversion_factor"), fmt_text)
    ws.write("A15", "Fator kgfm → kN·m", fmt_label); ws.write("B15", processed.summary.get("moment_conversion_factor"), fmt_text)
    ws.write("A17", "CONTAGENS", fmt_section)
    counts = [
        ("Combinações ELU", processed.summary.get("n_combinations_elu")),
        ("Combinações ELS", processed.summary.get("n_combinations_els")),
        ("Membros ELU/ELS", processed.summary.get("n_members")),
        ("Membros com ELS", processed.summary.get("n_members_els")),
        ("Nós de apoio", processed.summary.get("n_support_nodes")),
    ]
    for i, (label, val) in enumerate(counts, 18):
        ws.write(i - 1, 0, label, fmt_label)
        ws.write(i - 1, 1, val, fmt_text)
    ws.write(24, 0, "MAPEAMENTO DE CASOS", fmt_section)
    ws.write_row(25, 0, ["Caso Robot", "Grupo", "Carregamento", "Abreviação", "Incluir nas reações"], fmt_header)
    for r, item in enumerate(sorted(mappings, key=lambda x: x.case_id), 26):
        ws.write(r, 0, item.case_id, fmt_body)
        ws.write(r, 1, item.group if str(item.group).strip() else str(item.case_id), fmt_body)
        ws.write(r, 2, item.load_name, fmt_body_left)
        ws.write(r, 3, item.abbreviation, fmt_body)
        ws.write(r, 4, "Sim" if item.include_reactions else "Não", fmt_body)

    def add_table_sheet(
        name: str,
        title: str,
        headers: List[str],
        rows: List[List[object]],
        widths: List[float],
        left_cols=(),
        merge_specs: Sequence[tuple[int, Sequence[int]]] = (),
    ):
        sheet = wb.add_worksheet(name)
        sheet.hide_gridlines(2)
        last_col = len(headers) - 1
        sheet.merge_range(0, 0, 0, last_col, title, fmt_title)
        sheet.write_row(2, 0, headers, fmt_header)
        for ri, row in enumerate(rows, 3):
            for ci, value in enumerate(row):
                if isinstance(value, (int, float)) and ci not in (0, 1):
                    fmt = fmt_num
                else:
                    fmt = fmt_body_left if ci in left_cols else fmt_body
                if isinstance(value, str) and value.upper() == "OK":
                    fmt = fmt_ok
                elif isinstance(value, str) and value.upper() == "NÃO OK":
                    fmt = fmt_bad
                sheet.write(ri, ci, value, fmt)

        # Merge only categorical/identifier columns. Result columns (ratios,
        # forces, moments, cases and status) are deliberately never merged.
        # Data rows start at Excel row index 3 (zero-based).
        for col, parent_cols in merge_specs:
            for start, end, value in _merge_runs(rows, col, parent_cols):
                first_excel_row = 3 + start
                last_excel_row = 3 + end
                fmt = fmt_body_left if col in left_cols else fmt_body
                sheet.merge_range(first_excel_row, col, last_excel_row, col, value, fmt)

        for ci, width in enumerate(widths):
            sheet.set_column(ci, ci, width)
        sheet.freeze_panes(3, 0)
        sheet.repeat_rows(0, 2)
        # All sheets use portrait orientation, matching the Word deliverable.
        sheet.set_portrait()
        sheet.fit_to_pages(1, 0)
        sheet.set_margins(0.35, 0.35, 0.45, 0.45)
        return sheet

    add_table_sheet(
        "COMBINACOES_ELU", f"Combinações de Estado Limite Último (ELU) - {processed.summary['structure_name']}",
        ["Combinação", "Expressão para MC", "Definição Robot"], processed.combinations_elu,
        [20, 58, 45], left_cols=(0, 1, 2),
    )
    add_table_sheet(
        "COMBINACOES_ELS", f"Combinações de Estado Limite de Serviço (ELS) - {processed.summary['structure_name']}",
        ["Combinação", "Expressão para MC", "Definição Robot"], processed.combinations_els,
        [20, 58, 45], left_cols=(0, 1, 2),
    )
    add_table_sheet(
        "TABELA_8_ELU", f"Tabela 8 - Dimensionamento Eletrônico / Aproveitamento dos Membros - {processed.summary['structure_name']}",
        ["Membro", "Perfil", "Material", "Lay", "Laz", "Índice ELU", "Caso", "Status Tensão", "Status Esbeltez"],
        processed.table8,
        [25, 18, 21, 10, 10, 12, 20, 15, 17], left_cols=(0, 1, 2),
    )
    add_table_sheet(
        "TABELA_9_ELS", f"Tabela 9 - Deslocamento Eletrônico - {processed.summary['structure_name']}",
        ["Membro", "Perfil", "Ratio (uy)", "Caso (uy)", "Ratio (uz)", "Caso (uz)", "Ratio (vx)", "Caso (vx)", "Ratio (vy)", "Caso (vy)", "Status Flecha"],
        processed.table9,
        [24, 18, 11, 18, 11, 18, 11, 18, 11, 18, 14], left_cols=(0, 1),
    )
    add_table_sheet(
        "TABELA_10_REACOES", f"Tabela 10 - Reações nos apoios - {processed.summary['structure_name']}",
        ["Base/nó", "Grupo", "Carregamento", "Fx (kN)", "Fy (kN)", "Fz (kN)", "Mx (kN·m)", "My (kN·m)", "Mz (kN·m)"],
        processed.table10,
        [10, 10, 28, 12, 12, 12, 13, 13, 13], left_cols=(2,),
        merge_specs=((0, ()), (1, (0,))),
    )

    # VALIDAÇÕES
    ws = wb.add_worksheet("VALIDACOES")
    ws.hide_gridlines(2)
    ws.set_column("A:A", 14)
    ws.set_column("B:B", 110)
    ws.merge_range("A1:B1", "VALIDAÇÕES DO PROCESSAMENTO", fmt_title)
    ws.write_row("A3", ["Nível", "Mensagem"], fmt_header)
    for i, item in enumerate(processed.validations, 3):
        label = item.level.upper()
        fmt = fmt_ok if item.level == "success" else fmt_bad if item.level == "error" else fmt_body
        ws.write(i, 0, label, fmt)
        ws.write(i, 1, item.message, fmt_body_left)

    if include_raw_data:
        # Keep the raw sheets for traceability, but hide them from the normal workflow.
        ws = wb.add_worksheet("RAW_COMBINACOES")
        headers = ["Robot ID", "Nome", "Tipo análise", "Tipo combinação", "Natureza", "Definição"]
        ws.write_row(0, 0, headers, fmt_header)
        for i, c in enumerate(processed.raw_combinations, 1):
            ws.write_row(i, 0, [c.robot_id, c.name, c.analysis_type, c.combination_type, c.nature, c.definition], fmt_body_left)
        ws.hide()

        ws = wb.add_worksheet("RAW_REACOES")
        headers = ["Nó", "Caso", "É combinação", "FX", "FY", "FZ", "MX", "MY", "MZ", "Chave original"]
        ws.write_row(0, 0, headers, fmt_header)
        for i, r in enumerate(processed.raw_reactions, 1):
            ws.write_row(i, 0, [r.node, r.case_id, r.is_combination, r.fx, r.fy, r.fz, r.mx, r.my, r.mz, r.raw_key], fmt_body)
        ws.hide()

        ws = wb.add_worksheet("RAW_MEMBROS")
        headers = [
            "Membro", "Membro bruto", "Perfil", "Material", "Lay", "Laz", "Ratio ELU", "Caso ELU",
            "Ratio uy", "Caso uy", "Ratio uz", "Caso uz", "Ratio vx", "Caso vx", "Ratio vy", "Caso vy", "Nome automático"
        ]
        ws.write_row(0, 0, headers, fmt_header)
        for i, m in enumerate(processed.raw_members, 1):
            ws.write_row(i, 0, [
                m.member, m.raw_member, m.profile, m.material, m.lay, m.laz, m.ratio_elu, m.case_elu,
                m.ratio_uy, m.case_uy, m.ratio_uz, m.case_uz, m.ratio_vx, m.case_vx, m.ratio_vy, m.case_vy, m.auto_named
            ], fmt_body)
        ws.hide()

    wb.close()
    return bio.getvalue()


# ------------------------------- DOCX ----------------------------------------


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill.replace("#", ""))


def _set_cell_width(cell, width_cm: float) -> None:
    cell.width = Cm(width_cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), str(int(Cm(width_cm).twips)))


def _set_table_width_percent(table, percent: int = 100) -> None:
    """Equivalent to Word's 'AutoFit to Window' for the generated table."""
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "pct")
    tbl_w.set(qn("w:w"), str(int(percent * 50)))  # 5000 = 100%

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "autofit")


def _merge_word_vertical(
    table,
    rows: Sequence[Sequence[object]],
    col: int,
    parent_cols: Sequence[int],
    font_size: float,
    align,
) -> None:
    # +1 because Word table row 0 is the header.
    for start, end, value in _merge_runs(rows, col, parent_cols):
        top = table.cell(start + 1, col)
        bottom = table.cell(end + 1, col)
        merged = top.merge(bottom)
        _write_cell(merged, _format_value(value), font_size, align=align)


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _set_row_no_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def _write_cell(cell, value, size_pt: float, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("" if value is None else str(value))
    run.font.name = FONT_NAME
    run.font.size = Pt(size_pt)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color.replace("#", ""))
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _format_value(value) -> str:
    if isinstance(value, float):
        return f"{value:.2f}".replace(".", ",")
    return "" if value is None else str(value)


def _configure_section(section, landscape: bool):
    section.top_margin = Cm(1.0)
    section.bottom_margin = Cm(1.0)
    section.left_margin = Cm(1.0)
    section.right_margin = Cm(1.0)
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Cm(29.7)
        section.page_height = Cm(21.0)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)


def _add_title(doc: Document, title: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(title)
    run.font.name = FONT_NAME
    run.font.size = Pt(10)
    run.bold = True


def _add_word_table(
    doc: Document,
    title: str,
    headers: Sequence[str],
    rows: Sequence[Sequence[object]],
    widths_cm: Sequence[float],
    font_size: float,
    left_cols: Sequence[int] = (),
    merge_specs: Sequence[tuple[int, Sequence[int]]] = (),
) -> None:
    _add_title(doc, title)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    _set_table_width_percent(table, 100)
    fitted_widths = _scale_widths(widths_cm, 18.8)

    hdr = table.rows[0]
    _set_repeat_table_header(hdr)
    _set_row_no_split(hdr)
    for j, h in enumerate(headers):
        cell = hdr.cells[j]
        _set_cell_width(cell, fitted_widths[j])
        _set_cell_shading(cell, DARK)
        _write_cell(cell, h, font_size, bold=True, color=PINK)

    for row_data in rows:
        row = table.add_row()
        _set_row_no_split(row)
        for j, value in enumerate(row_data):
            cell = row.cells[j]
            _set_cell_width(cell, fitted_widths[j])
            text = _format_value(value)
            align = WD_ALIGN_PARAGRAPH.LEFT if j in left_cols else WD_ALIGN_PARAGRAPH.CENTER
            _write_cell(cell, text, font_size, align=align)
            if isinstance(value, str) and value.upper() == "OK":
                _set_cell_shading(cell, GREEN)
                for run in cell.paragraphs[0].runs:
                    run.font.color.rgb = RGBColor.from_string(GREEN_TXT.replace("#", ""))
                    run.bold = True
            elif isinstance(value, str) and value.upper() == "NÃO OK":
                _set_cell_shading(cell, RED)
                for run in cell.paragraphs[0].runs:
                    run.font.color.rgb = RGBColor.from_string(RED_TXT.replace("#", ""))
                    run.bold = True

    for col, parent_cols in merge_specs:
        align = WD_ALIGN_PARAGRAPH.LEFT if col in left_cols else WD_ALIGN_PARAGRAPH.CENTER
        _merge_word_vertical(table, rows, col, parent_cols, font_size, align)

    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def _new_section(doc: Document, landscape: bool = False):
    # Kept as a helper for pagination, but all generated pages are portrait.
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    _configure_section(section, landscape=False)
    return section


def build_word(processed: ProcessedData, mappings: Sequence[CaseMapping], input_names: Dict[str, str]) -> bytes:
    doc = Document()
    _configure_section(doc.sections[0], landscape=False)

    # Base style
    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(9)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("TABELAS DE RESULTADOS - ROBOT STRUCTURAL ANALYSIS")
    r.font.name = FONT_NAME
    r.font.size = Pt(13)
    r.bold = True

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = subtitle.add_run(str(processed.summary.get("structure_name", "")))
    rr.font.name = FONT_NAME
    rr.font.size = Pt(11)
    rr.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(
        f"Arquivo gerado automaticamente. Combinações ELU: {processed.summary['n_combinations_elu']} | "
        f"ELS: {processed.summary['n_combinations_els']} | Membros: {processed.summary['n_members']} | "
        f"Nós de apoio: {processed.summary['n_support_nodes']}."
    )
    run.font.name = FONT_NAME
    run.font.size = Pt(9)

    # Validation summary (short)
    p = doc.add_paragraph()
    r = p.add_run("Validação: ")
    r.font.name = FONT_NAME; r.font.size = Pt(9); r.bold = True
    r2 = p.add_run(_status_level(processed))
    r2.font.name = FONT_NAME; r2.font.size = Pt(9)

    _new_section(doc, landscape=False)
    _add_word_table(
        doc,
        f"Combinações de Estado Limite Último (ELU) - {processed.summary['structure_name']}",
        ["Combinação", "Expressão para MC"],
        [[r[0], r[1]] for r in processed.combinations_elu],
        [3.7, 14.5],
        8.0,
        left_cols=(0, 1),
    )

    _new_section(doc, landscape=False)
    _add_word_table(
        doc,
        f"Combinações de Estado Limite de Serviço (ELS) - {processed.summary['structure_name']}",
        ["Combinação", "Expressão para MC"],
        [[r[0], r[1]] for r in processed.combinations_els],
        [3.7, 14.5],
        8.0,
        left_cols=(0, 1),
    )

    _new_section(doc, landscape=False)
    _add_word_table(
        doc,
        f"Tabela 8 - Dimensionamento Eletrônico / Aproveitamento dos Membros - {processed.summary['structure_name']}",
        ["Membro", "Perfil", "Material", "Lay", "Laz", "Índice ELU", "Caso", "Status Tensão", "Status Esbeltez"],
        processed.table8,
        [4.0, 3.0, 3.3, 1.5, 1.5, 1.8, 3.0, 2.4, 2.6],
        6.3,
        left_cols=(0, 1, 2),
    )

    _new_section(doc, landscape=False)
    _add_word_table(
        doc,
        f"Tabela 9 - Deslocamento Eletrônico - {processed.summary['structure_name']}",
        ["Membro", "Perfil", "Ratio (uy)", "Caso (uy)", "Ratio (uz)", "Caso (uz)", "Ratio (vx)", "Caso (vx)", "Ratio (vy)", "Caso (vy)", "Status Flecha"],
        processed.table9,
        [3.2, 2.4, 1.4, 2.3, 1.4, 2.3, 1.4, 2.3, 1.4, 2.3, 1.8],
        5.5,
        left_cols=(0, 1),
    )

    _new_section(doc, landscape=False)
    _add_word_table(
        doc,
        f"Tabela 10 - Reações nos apoios - {processed.summary['structure_name']}",
        ["Base/nó", "Grupo", "Carregamento", "Fx (kN)", "Fy (kN)", "Fz (kN)", "Mx (kN·m)", "My (kN·m)", "Mz (kN·m)"],
        processed.table10,
        [1.6, 1.5, 5.0, 2.2, 2.2, 2.2, 2.5, 2.5, 2.5],
        6.0,
        left_cols=(2,),
        merge_specs=((0, ()), (1, (0,))),
    )

    bio = BytesIO()
    doc.save(bio)
    return bio.getvalue()
